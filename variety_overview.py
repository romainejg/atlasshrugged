from __future__ import annotations

import base64
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import winreg
from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw


IMAGE_SHEET = "Output4 - Photo Audit"
MODEL_WORKFLOW = "leafy"
GRID_SQUARE_CM = 1.27
_ROBOFLOW_WORKSPACE = None


def key_name(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def roboflow_key() -> str:
    value = os.environ.get("ROBOFLOW_API_KEY", "").strip()
    if value:
        return value
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as registry:
            value = str(winreg.QueryValueEx(registry, "ROBOFLOW_API_KEY")[0]).strip()
    except OSError:
        value = ""
    if value:
        os.environ["ROBOFLOW_API_KEY"] = value
    return value


def save_roboflow_key(value: str) -> None:
    with winreg.OpenKey(
        winreg.HKEY_CURRENT_USER, "Environment", 0, winreg.KEY_SET_VALUE
    ) as registry:
        winreg.SetValueEx(registry, "ROBOFLOW_API_KEY", 0, winreg.REG_SZ, value)
    os.environ["ROBOFLOW_API_KEY"] = value


def grid_spacing_pixels(image: Image.Image) -> float | None:
    rgb = np.asarray(image.convert("RGB"), dtype=np.float32)
    height, width, _ = rgb.shape
    gray = rgb.mean(axis=2)
    region = gray[int(height * 0.12):int(height * 0.62), int(width * 0.06):int(width * 0.94)]
    projection = (region < 120).mean(axis=0)
    candidate = np.where(projection >= max(0.01, float(np.percentile(projection, 94))))[0]
    if candidate.size < 5:
        return None
    groups = []
    start = previous = int(candidate[0])
    for value in candidate[1:]:
        value = int(value)
        if value > previous + 2:
            groups.append((start + previous) / 2)
            start = value
        previous = value
    groups.append((start + previous) / 2)
    gaps = np.diff(groups)
    valid = gaps[(gaps > width * 0.018) & (gaps < width * 0.08)]
    return float(np.median(valid)) if valid.size else None


def workflow_predictions(image_bytes: bytes, api_key: str) -> list[dict]:
    global _ROBOFLOW_WORKSPACE
    workspace = _ROBOFLOW_WORKSPACE
    if not workspace:
        auth = Request(
            f"https://api.roboflow.com/?api_key={api_key}",
            headers={"Accept": "application/json"},
        )
        with urlopen(auth, timeout=60) as response:
            workspace = json.loads(response.read().decode("utf-8")).get("workspace")
        _ROBOFLOW_WORKSPACE = workspace
    if not workspace:
        raise RuntimeError("Roboflow did not return a workspace.")
    workflow_id = os.environ.get("ROBOFLOW_WORKFLOW_ID", MODEL_WORKFLOW)
    payload = {
        "api_key": api_key,
        "inputs": {
            "image": {
                "type": "base64",
                "value": base64.b64encode(image_bytes).decode("ascii"),
            }
        },
        "use_cache": True,
    }
    request = Request(
        f"https://serverless.roboflow.com/{workspace}/workflows/{workflow_id}",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json", "Accept": "application/json"},
        method="POST",
    )
    with urlopen(request, timeout=120) as response:
        result = json.loads(response.read().decode("utf-8"))

    def find(value):
        if isinstance(value, dict):
            if isinstance(value.get("predictions"), list):
                return value["predictions"]
            for nested in value.values():
                found = find(nested)
                if found is not None:
                    return found
        if isinstance(value, list):
            for nested in value:
                found = find(nested)
                if found is not None:
                    return found
        return None

    predictions = find(result)
    if predictions is None:
        raise RuntimeError("Workflow Leafy returned no instance predictions.")
    return predictions


def cutout_largest_leaf(image_bytes: bytes, destination: Path, api_key: str) -> dict | None:
    source = Image.open(BytesIO(image_bytes)).convert("RGBA")
    predictions = workflow_predictions(image_bytes, api_key)
    candidates = [p for p in predictions if p.get("points")]
    if not candidates:
        return None
    prediction = max(
        candidates,
        key=lambda p: max(pt["y"] for pt in p["points"]) - min(pt["y"] for pt in p["points"]),
    )
    points = [(float(point["x"]), float(point["y"])) for point in prediction["points"]]
    left = max(0, int(math.floor(min(x for x, _ in points))))
    top = max(0, int(math.floor(min(y for _, y in points))))
    right = min(source.width, int(math.ceil(max(x for x, _ in points))) + 1)
    bottom = min(source.height, int(math.ceil(max(y for _, y in points))) + 1)
    mask = Image.new("L", source.size, 0)
    ImageDraw.Draw(mask).polygon(points, fill=255)
    rgba = Image.new("RGBA", source.size, (255, 255, 255, 0))
    rgba.paste(source, (0, 0), mask)
    cropped = rgba.crop((left, top, right, bottom))
    destination.parent.mkdir(parents=True, exist_ok=True)
    cropped.save(destination)
    spacing = grid_spacing_pixels(source)
    height_cm = (bottom - top) / spacing * GRID_SQUARE_CM if spacing else None
    width_cm = (right - left) / spacing * GRID_SQUARE_CM if spacing else None
    return {
        "path": str(destination.resolve()),
        "height_cm": round(height_cm, 1) if height_cm else None,
        "width_cm": round(width_cm, 1) if width_cm else None,
        "pixel_width": cropped.width,
        "pixel_height": cropped.height,
    }


def extract_workbook(workbook_path: Path, asset_dir: Path, progress=None, only_varieties=None) -> dict:
    workbook = load_workbook(workbook_path, data_only=True)
    environment = {}
    environment_by_dtm = {}
    if "Output5 - Environment" in workbook.sheetnames:
        sheet = workbook["Output5 - Environment"]
        headers = {key_name(cell.value): cell.column for cell in sheet[1] if cell.value}
        rows = [row for row in range(2, sheet.max_row + 1) if isinstance(sheet.cell(row, 1).value, (int, float))]
        dli_col = next((column for key, column in headers.items() if "daylightintegral" in key), None)
        temp_col = next((column for key, column in headers.items() if "temperaturec24haranet" in key), None)
        if temp_col is None:
            temp_col = next((column for key, column in headers.items() if "temperature" in key and "average" in key), None)
        rh_col = next((column for key, column in headers.items() if key.startswith("humidity") and "24haranet" in key), None)
        harvest_col = headers.get(key_name("Harvest Date"))
        for row in rows:
            dtm = int(sheet.cell(row, 1).value)
            environment_by_dtm[dtm] = {
                "dtm": dtm,
                "dli": sheet.cell(row, dli_col).value if dli_col else None,
                "temperature": sheet.cell(row, temp_col).value if temp_col else None,
                "rh": sheet.cell(row, rh_col).value if rh_col else None,
                "harvest_date": (
                    sheet.cell(row, harvest_col).value.strftime("%Y-%m-%d")
                    if harvest_col and hasattr(sheet.cell(row, harvest_col).value, "strftime")
                    else sheet.cell(row, harvest_col).value if harvest_col else None
                ),
            }
        if rows:
            environment = environment_by_dtm[max(environment_by_dtm)]

    output1 = workbook["Output1 - Excel"]
    headers1 = {key_name(cell.value): cell.column for cell in output1[1] if cell.value}
    variety_col = headers1[key_name("Variety")]
    cycle_col = headers1[key_name("Cycle legnth")]
    yield_col = headers1[key_name("Annualized Yield (kg/m2/yr)")]
    trial_col = headers1[key_name("Trial")]
    leaf_size_col = headers1.get(key_name("Leaf Size (cm)"))
    gmol_col = headers1.get(key_name("Light Use Efficiency (g/mol)")) or headers1.get(key_name("g/mol"))
    varieties: dict[str, dict] = {}
    for row in range(2, output1.max_row + 1):
        name = output1.cell(row, variety_col).value
        cycle = output1.cell(row, cycle_col).value
        annual_yield = output1.cell(row, yield_col).value
        if not name or not isinstance(cycle, (int, float)):
            continue
        yield_was_normalized = False
        if isinstance(annual_yield, (int, float)) and annual_yield > 1000:
            # Older brassica analyses treated gram weights as kilograms. Keep
            # existing-file brochure runs safe by correcting that unmistakable
            # 1000x unit error before charting it as kg/m2/yr.
            annual_yield = round(annual_yield / 1000, 2)
            yield_was_normalized = True
        item = varieties.setdefault(
            key_name(name),
            {"name": str(name), "trial": output1.cell(row, trial_col).value, "yield": [], "leaves": []},
        )
        dtm = int(cycle)
        dli = (environment_by_dtm.get(dtm) or {}).get("dli")
        gmol = output1.cell(row, gmol_col).value if gmol_col else None
        if (
            yield_was_normalized
            or not isinstance(gmol, (int, float))
        ) and isinstance(annual_yield, (int, float)) and isinstance(dli, (int, float)) and dli:
            gmol = round(annual_yield * 1000 / (365 * dli), 2)
        item["yield"].append({
            "dtm": dtm,
            "value": annual_yield if isinstance(annual_yield, (int, float)) else None,
            "gmol": gmol if isinstance(gmol, (int, float)) else None,
            "height_cm": output1.cell(row, leaf_size_col).value if leaf_size_col else None,
        })

    if "Output3 - Airtable" in workbook.sheetnames:
        sheet = workbook["Output3 - Airtable"]
        headers = {key_name(cell.value): cell.column for cell in sheet[2] if cell.value}
        for row in range(3, sheet.max_row + 1):
            name = sheet.cell(row, headers.get("variety", 3)).value
            item = varieties.get(key_name(name))
            if item:
                item["strengths"] = sheet.cell(row, headers.get("pros", 5)).value
                item["weaknesses"] = sheet.cell(row, headers.get("cons", 6)).value

    api_key = roboflow_key()
    audit = workbook[IMAGE_SHEET]
    labels = {}
    for row in range(1, audit.max_row + 1):
        value = audit.cell(row, 2).value
        match = re.match(r"(.+?)\s*[-–—]\s*DTM\s+(\d+)\s*[-–—]\s*Original", str(value or ""), re.I)
        if match:
            labels[row + 1] = (match.group(1).strip(), int(match.group(2)))
    originals = [
        image for image in audit._images
        if image.anchor._from.col + 1 == 2 and image.anchor._from.row + 1 in labels
    ]
    total = len(originals)
    for index, image in enumerate(originals, start=1):
        anchor_row = image.anchor._from.row + 1
        name, dtm = labels[anchor_row]
        if only_varieties and key_name(name) not in {key_name(value) for value in only_varieties}:
            continue
        item = varieties.get(key_name(name))
        if item is None:
            continue
        if progress:
            progress(f"Removing background: {name}, DTM {dtm} ({index}/{total})")
        safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", name)
        destination = asset_dir / f"{safe_name}_{dtm}_leaf.png"
        if destination.exists():
            cached = Image.open(destination)
            result = {
                "path": str(destination.resolve()), "height_cm": None, "width_cm": None,
                "pixel_width": cached.width, "pixel_height": cached.height,
            }
        else:
            if not api_key:
                raise RuntimeError("Roboflow API key is not configured and no saved render exists for this photo.")
            result = cutout_largest_leaf(image._data(), destination, api_key)
        if result:
            result["dtm"] = dtm
            matching_point = next((point for point in item["yield"] if point["dtm"] == dtm), None)
            if matching_point and not result.get("height_cm"):
                result["height_cm"] = matching_point.get("height_cm")
            if (
                result.get("height_cm")
                and not result.get("width_cm")
                and result.get("pixel_height")
                and result.get("pixel_width")
            ):
                result["width_cm"] = round(
                    result["height_cm"] * result["pixel_width"] / result["pixel_height"],
                    1,
                )
            result["selection"] = "largest Roboflow leaf by segmented height"
            item["leaves"].append(result)

    for item in varieties.values():
        item["yield"].sort(key=lambda point: point["dtm"])
        item["leaves"].sort(key=lambda leaf: leaf["dtm"])
        item.setdefault("strengths", None)
        item.setdefault("weaknesses", None)
    return {
        "source": str(workbook_path.resolve()),
        "trial": next((item["trial"] for item in varieties.values()), workbook_path.stem),
        "environment": environment,
        "varieties": list(varieties.values()),
    }


def find_node() -> Path:
    found = shutil.which("node")
    if found:
        return Path(found)
    candidates = list(Path.home().glob(".cache/codex-runtimes/*/dependencies/node/bin/node.exe"))
    if candidates:
        return candidates[0]
    raise RuntimeError("Node.js was not found. Install Node.js to create PowerPoint files.")


def create_overview(workbook_path: Path, output_folder: Path, progress=None) -> Path:
    app_dir = Path(__file__).resolve().parent
    output_folder.mkdir(parents=True, exist_ok=True)
    work_dir = Path(tempfile.mkdtemp(prefix="variety-overview-", dir=output_folder))
    try:
        # Keep Roboflow-rendered cutouts alongside the analysis output.  Existing
        # files are intentionally reused in extract_workbook rather than re-run.
        data = extract_workbook(workbook_path, output_folder / "Variety Overview Photos", progress)
        data_path = work_dir / "data.json"
        data_path.write_text(json.dumps(data, indent=2), encoding="utf-8")
        output = output_folder / f"Variety Overview - {data['trial']}.pptx"
        environment = os.environ.copy()
        environment["VARIETY_DATA"] = str(data_path)
        environment["VARIETY_OUTPUT"] = str(output.resolve())
        if progress:
            progress("Building PowerPoint…")
        builder = app_dir / ".pptx-build" / "build-variety-overview.ps1"
        if not builder.is_file():
            raise RuntimeError(f"PowerPoint builder is missing: {builder}")
        subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(builder)],
            cwd=app_dir,
            env=environment,
            check=True,
            capture_output=True,
            text=True,
        )
        return output
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
