"""
Deck builder for Biophi Suite.

This module is the server-compatible (Linux/macOS/Windows) replacement for
``build_brochure.ps1``, which relied on Windows COM automation.  It uses the
``python-pptx`` library to generate a branded PowerPoint presentation from
analysis results and the ``template.pptx`` slide master.

Usage (CLI):
    python build_deck.py --output /path/to/output.pptx

Usage (API):
    from build_deck import build_deck
    output_path = build_deck(analysis_results, output_path="brochure.pptx")
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

_TEMPLATE_PATH = Path(__file__).parent.parent / "template.pptx"
_DEFAULT_OUTPUT = Path(__file__).parent.parent / "output" / "brochure.pptx"
_IMAGE_CACHE_DIR = Path(__file__).parent.parent / "Assets" / "image_cache"


def build_deck(
    analysis_results: List[Dict[str, Any]],
    output_path: Optional[str] = None,
    template_path: Optional[str] = None,
) -> Path:
    """
    Generate a PowerPoint presentation from *analysis_results*.

    Args:
        analysis_results: List of result dicts produced by ``analysis.analyse_all``.
        output_path:      Destination ``.pptx`` file.  Defaults to
                          ``../output/brochure.pptx`` relative to this file.
        template_path:    Path to a branded ``.pptx`` template.  Defaults to
                          ``../template.pptx``.

    Returns:
        The resolved path of the generated file.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise ImportError(
            "python-pptx is required to build decks. "
            "Install it with:  pip install python-pptx"
        ) from exc

    tmpl = Path(template_path) if template_path else _TEMPLATE_PATH
    out = Path(output_path) if output_path else _DEFAULT_OUTPUT
    out.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(tmpl)) if tmpl.exists() else Presentation()

    # Use the first two slide layouts: title slide and a content slide.
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]

    for result in analysis_results:
        fields: Dict[str, Any] = result.get("fields", {})
        predictions: List[Dict[str, Any]] = result.get("predictions", [])

        # --- Title slide for this record ---
        title_text = fields.get("Name") or fields.get("Title") or result.get("record_id", "")
        subtitle_text = fields.get("Description") or fields.get("Notes") or ""

        slide = prs.slides.add_slide(title_layout)
        _set_placeholder_text(slide, 0, str(title_text))
        _set_placeholder_text(slide, 1, str(subtitle_text))

        # --- Content slide: one per predicted image ---
        for pred in predictions:
            filename: str = pred.get("filename", "")
            image_path = _IMAGE_CACHE_DIR / filename
            result_data: Dict[str, Any] = pred.get("result", {})

            content_slide = prs.slides.add_slide(content_layout)
            _set_placeholder_text(content_slide, 0, filename)

            summary_lines = _summarise_predictions(result_data)
            _set_placeholder_text(content_slide, 1, "\n".join(summary_lines))

            if image_path.exists():
                _add_image(content_slide, str(image_path))

    prs.save(str(out))
    logger.info("Deck saved to %s", out)
    return out


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _set_placeholder_text(slide: Any, idx: int, text: str) -> None:
    """Set text on placeholder *idx* if it exists; silently skip otherwise."""
    try:
        slide.placeholders[idx].text = text
    except (KeyError, IndexError):
        pass


def _add_image(slide: Any, image_path: str) -> None:
    """Add an image to the slide, anchored at the top-right quadrant."""
    try:
        from pptx.util import Inches
        slide.shapes.add_picture(image_path, Inches(5), Inches(1.5), width=Inches(4))
    except Exception:
        logger.warning("Could not add image %s to slide.", image_path)


def _summarise_predictions(result: Dict[str, Any]) -> List[str]:
    """Convert a Roboflow result dict into human-readable bullet lines."""
    lines: List[str] = []
    for pred in result.get("predictions", []):
        label = pred.get("class", "object")
        conf = pred.get("confidence", 0)
        lines.append(f"• {label}: {conf:.0%} confidence")
    if not lines:
        lines.append("No detections above threshold.")
    return lines


# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

def _parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a Biophi brochure deck.")
    parser.add_argument(
        "--output", default=str(_DEFAULT_OUTPUT), help="Output .pptx file path."
    )
    parser.add_argument(
        "--template", default=str(_TEMPLATE_PATH), help="Branded template .pptx path."
    )
    parser.add_argument(
        "--force-refresh", action="store_true", help="Re-fetch Airtable records."
    )
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    args = _parse_args(argv)

    sys.path.insert(0, str(Path(__file__).parent))
    from analysis import analyse_all

    results = analyse_all(
        model_id=_require_model_id(),
        force_airtable_refresh=args.force_refresh,
    )
    out = build_deck(results, output_path=args.output, template_path=args.template)
    print("Deck written to:", out)
    return 0


def _require_model_id() -> str:
    import os

    model_id = os.environ.get("ROBOFLOW_MODEL_ID", "")
    if not model_id:
        raise EnvironmentError(
            "ROBOFLOW_MODEL_ID environment variable is not set. "
            "Set it to your Roboflow model in the form 'project/version'."
        )
    return model_id


if __name__ == "__main__":
    sys.exit(main())
